"""In-image checks used by the release runtime regression harness.

This file deliberately generates all fixtures at runtime.  It is bind-mounted
into the image built from this repository, so the checks exercise the exact
LibreOffice, Tesseract, Python packages, Alembic migrations and ORM code that
the production image contains.  It must never be pointed at a production
database or uploads directory.
"""
from __future__ import annotations

import argparse
import hashlib
import os
import subprocess
import sys
import tempfile
import time
from datetime import timedelta
from pathlib import Path


PROJECT_ROOT = Path("/app")
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))


def _assert(condition: bool, message: str) -> None:
    if not condition:
        raise RuntimeError(message)


def _run_alembic(*arguments: str) -> None:
    """Run a migration without printing the database URL on failure."""

    completed = subprocess.run(
        ["alembic", *arguments],
        cwd=PROJECT_ROOT,
        text=True,
        capture_output=True,
        check=False,
    )
    if completed.returncode != 0:
        # Alembic normally does not print the configured connection URL, but
        # do not relay child output here because this harness is intentionally
        # safe to use with an ephemeral generated database password.
        raise RuntimeError(f"alembic_{arguments[0]}_failed")


def _make_pdf(path: Path, marker: str, *, font_size: float = 24) -> None:
    import fitz

    document = fitz.open()
    try:
        page = document.new_page(width=612, height=792)
        page.insert_text((54, 120), marker, fontsize=font_size, fontname="helv")
        document.save(path)
    finally:
        document.close()


def _make_image(path: Path, marker: str) -> None:
    """Render real raster text without adding a new image-library dependency."""

    import fitz

    with tempfile.TemporaryDirectory() as temporary:
        source_pdf = Path(temporary) / "image-source.pdf"
        _make_pdf(source_pdf, marker, font_size=34)
        document = fitz.open(source_pdf)
        try:
            pixmap = document.load_page(0).get_pixmap(matrix=fitz.Matrix(3, 3), alpha=False)
            if path.suffix.lower() == ".png":
                pixmap.save(path)
            else:
                # PyMuPDF writes a real JPEG byte stream here.  Giving this
                # path a .jpg suffix makes the application's actual JPG route
                # select Tesseract rather than a PDF fallback.
                path.write_bytes(pixmap.tobytes("jpeg"))
        finally:
            document.close()


def _assert_marker(result: object, *, parser_fragment: str, marker_words: tuple[str, ...]) -> None:
    parser_version = str(getattr(result, "parser_version"))
    raw_text = str(getattr(result, "raw_text"))
    _assert(parser_fragment in parser_version, f"unexpected_parser:{parser_version}")
    normalized = raw_text.upper()
    _assert(
        all(word in normalized for word in marker_words),
        f"synthetic_marker_missing:{parser_fragment}",
    )
    _assert(int(getattr(result, "source_page_count")) >= 1, "source_page_count_missing")
    _assert(int(getattr(result, "parsed_page_count")) >= 1, "parsed_page_count_missing")


def run_document_regression() -> None:
    """Exercise PDF, DOCX, XLSX, PNG, JPG and HTML in the production image."""

    from docx import Document
    from openpyxl import Workbook

    from app.services.document_text_extraction import extract_document_text

    with tempfile.TemporaryDirectory(prefix="greatsell-document-regression-") as temporary:
        fixtures = Path(temporary)
        pdf_path = fixtures / "synthetic-resume.pdf"
        docx_path = fixtures / "synthetic-resume.docx"
        xlsx_path = fixtures / "synthetic-resume.xlsx"
        png_path = fixtures / "synthetic-resume.png"
        jpg_path = fixtures / "synthetic-resume.jpg"
        html_path = fixtures / "synthetic-resume.html"

        _make_pdf(pdf_path, "SYNTHETIC PDF RESUME MARKER")

        document = Document()
        document.add_heading("Synthetic resume", level=1)
        document.add_paragraph("SYNTHETIC DOCX RESUME MARKER")
        document.save(docx_path)

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Resume"
        worksheet.append(["SYNTHETIC", "XLSX", "RESUME", "MARKER"])
        workbook.save(xlsx_path)
        workbook.close()

        _make_image(png_path, "SYNTHETIC PNG RESUME MARKER")
        _make_image(jpg_path, "SYNTHETIC JPG RESUME MARKER")

        html_path.write_text(
            "<html><body><script>window.RUNTIME_SCRIPT_MARKER = true;</script>"
            "<main>SYNTHETIC HTML RESUME MARKER</main></body></html>",
            encoding="utf-8",
        )

        extraction_options = {
            "min_text_chars_per_page": 1,
            "ocr_sparse_text_chars_per_page": 1,
            "tencent_ocr_config": None,
        }
        pdf_result = extract_document_text(pdf_path, **extraction_options)
        _assert_marker(
            pdf_result,
            parser_fragment="pypdf-",
            marker_words=("SYNTHETIC", "PDF", "RESUME"),
        )

        docx_result = extract_document_text(docx_path, **extraction_options)
        _assert_marker(
            docx_result,
            parser_fragment="pypdf-",
            marker_words=("SYNTHETIC", "DOCX", "RESUME"),
        )

        xlsx_result = extract_document_text(xlsx_path, **extraction_options)
        _assert_marker(
            xlsx_result,
            parser_fragment="openpyxl",
            marker_words=("SYNTHETIC", "XLSX", "RESUME"),
        )

        png_result = extract_document_text(png_path, **extraction_options)
        _assert_marker(
            png_result,
            parser_fragment="tesseract",
            marker_words=("SYNTHETIC", "PNG", "RESUME"),
        )

        jpg_result = extract_document_text(jpg_path, **extraction_options)
        _assert_marker(
            jpg_result,
            parser_fragment="tesseract",
            marker_words=("SYNTHETIC", "JPG", "RESUME"),
        )

        html_result = extract_document_text(html_path, **extraction_options)
        _assert_marker(
            html_result,
            parser_fragment="beautifulsoup4",
            marker_words=("SYNTHETIC", "HTML", "RESUME"),
        )
        _assert(
            "RUNTIME_SCRIPT_MARKER" not in html_result.raw_text,
            "html_script_was_not_removed_before_extraction",
        )

    print("runtime-document-regression: passed")


def _database_url() -> str:
    value = os.getenv("RESUME_V3_DATABASE_URL", "").strip()
    if not value.startswith("postgresql+"):
        raise RuntimeError("release_regression_requires_postgresql_url")
    return value


def _uploads_dir() -> Path:
    raw_path = os.getenv("RELEASE_REGRESSION_UPLOADS_DIR", "").strip()
    if not raw_path:
        raise RuntimeError("release_regression_uploads_dir_required")
    path = Path(raw_path)
    path.mkdir(parents=True, exist_ok=True)
    return path


def _expected_alembic_head() -> str:
    from alembic.config import Config
    from alembic.script import ScriptDirectory

    return str(ScriptDirectory.from_config(Config(str(PROJECT_ROOT / "alembic.ini"))).get_current_head())


def _assert_current_head(database: object, expected_head: str) -> None:
    from sqlalchemy import text

    session_factory = getattr(database, "session_factory")
    with session_factory() as session:
        actual_head = session.scalar(text("SELECT version_num FROM alembic_version"))
    _assert(actual_head == expected_head, "alembic_head_mismatch")


def _settings(database_url: str, uploads_dir: Path) -> object:
    from app.config import AppSettings

    return AppSettings(
        project_dir=PROJECT_ROOT,
        data_dir=uploads_dir.parent,
        upload_dir=uploads_dir,
        database_url=database_url,
        auto_create_schema=False,
        seed_registry_on_startup=False,
    )


def run_database_seed() -> None:
    """Migrate from the first revision and write synthetic backup material."""

    from datetime import datetime, timezone

    from app.database import Database
    from app.models import Candidate, MailboxBackgroundJob, MailboxConfig, Organization, Resume
    from app.tenant_scope import clear_organization_context, set_organization_context

    database_url = _database_url()
    uploads_dir = _uploads_dir()
    _run_alembic("upgrade", "20260716_0001")
    database = Database(database_url)
    try:
        _assert_current_head(database, "20260716_0001")
    finally:
        database.dispose()

    _run_alembic("upgrade", "head")
    database = Database(database_url)
    expected_head = _expected_alembic_head()
    _assert_current_head(database, expected_head)

    payload = b"GreatSell synthetic recovery original. No candidate data.\n"
    digest = hashlib.sha256(payload).hexdigest()
    now = datetime.now(timezone.utc)
    try:
        with database.session_factory() as session:
            primary_organization = Organization(name="Synthetic recovery workspace")
            secondary_organization = Organization(name="Synthetic isolated workspace")
            session.add_all((primary_organization, secondary_organization))
            session.flush()

            set_organization_context(session, primary_organization.id)
            try:
                candidate = Candidate(display_name="Synthetic candidate")
                session.add(candidate)
                session.flush()
                storage_key = f"{primary_organization.id}/synthetic-recovery-resume.pdf"
                original_path = uploads_dir / storage_key
                original_path.parent.mkdir(parents=True, exist_ok=True)
                original_path.write_bytes(payload)
                resume = Resume(
                    candidate_id=candidate.id,
                    original_filename="synthetic-recovery-resume.pdf",
                    storage_key=storage_key,
                    sha256=digest,
                    source_page_count=1,
                    parsed_page_count=1,
                    extraction_status="text_ready",
                    quality_flags=[],
                    parser_version="release-regression",
                    is_active=False,
                    facts_version=0,
                    raw_text="SYNTHETIC RECOVERY MARKER",
                )
                mailbox = MailboxConfig(
                    display_name="Synthetic recovery mailbox",
                    display_name_key="synthetic recovery mailbox",
                    imap_host="imap.invalid.test",
                    imap_port=993,
                    email_address="synthetic-recovery@invalid.test",
                    mailbox="INBOX",
                    encrypted_password="synthetic-not-a-secret",
                    enabled=True,
                )
                session.add_all((resume, mailbox))
                session.flush()
                expired_job = MailboxBackgroundJob(
                    mailbox_config_id=mailbox.id,
                    job_kind="sync",
                    trigger_type="manual",
                    status="running",
                    attempt_count=1,
                    max_attempts=3,
                    lease_owner="crashed-worker",
                    lease_expires_at=now - timedelta(minutes=5),
                    requested_at=now - timedelta(minutes=10),
                    started_at=now - timedelta(minutes=6),
                )
                session.add(expired_job)
                # Flush before changing workspace context.  The ORM tenant
                # guard intentionally rejects a pending row from workspace A
                # if it is accidentally flushed while workspace B is active.
                session.flush()
            finally:
                clear_organization_context(session)

            set_organization_context(session, secondary_organization.id)
            try:
                secondary_mailbox = MailboxConfig(
                    display_name="Synthetic isolated mailbox",
                    display_name_key="synthetic isolated mailbox",
                    imap_host="imap.invalid.test",
                    imap_port=993,
                    email_address="synthetic-isolated@invalid.test",
                    mailbox="INBOX",
                    encrypted_password="synthetic-not-a-secret",
                    enabled=True,
                )
                session.add(secondary_mailbox)
                session.flush()
                untouched_job = MailboxBackgroundJob(
                    mailbox_config_id=secondary_mailbox.id,
                    job_kind="sync",
                    trigger_type="manual",
                    status="queued",
                    attempt_count=0,
                    max_attempts=3,
                    next_attempt_at=now + timedelta(days=1),
                    requested_at=now,
                )
                session.add(untouched_job)
                session.flush()
            finally:
                clear_organization_context(session)
            session.commit()
    finally:
        database.dispose()

    print("runtime-postgres-seed: passed")


def run_database_verify() -> None:
    """Verify restored data and run the worker's real expired-lease path."""

    from datetime import datetime, timezone

    from sqlalchemy import select

    from app.database import Database
    from app.models import MailboxBackgroundJob, Organization, Resume
    from app.services import mailbox_background_job_service
    from app.tenant_scope import clear_organization_context, set_organization_context

    database_url = _database_url()
    uploads_dir = _uploads_dir()
    database = Database(database_url)
    expected_head = _expected_alembic_head()
    _assert_current_head(database, expected_head)
    settings = _settings(database_url, uploads_dir)

    try:
        with database.session_factory() as session:
            primary = session.scalar(
                select(Organization).where(Organization.name == "Synthetic recovery workspace")
            )
            secondary = session.scalar(
                select(Organization).where(Organization.name == "Synthetic isolated workspace")
            )
            _assert(primary is not None and secondary is not None, "restored_workspaces_missing")

            set_organization_context(session, primary.id)
            try:
                resume = session.scalar(
                    select(Resume).where(
                        Resume.original_filename == "synthetic-recovery-resume.pdf"
                    )
                )
                expired_job = session.scalar(
                    select(MailboxBackgroundJob).where(
                        MailboxBackgroundJob.status == "running",
                        MailboxBackgroundJob.lease_owner == "crashed-worker",
                    )
                )
            finally:
                clear_organization_context(session)
            _assert(resume is not None, "restored_resume_missing")
            _assert(expired_job is not None, "restored_expired_job_missing")
            _assert(expired_job.organization_id == primary.id, "restored_job_workspace_mismatch")
            original_path = uploads_dir / resume.storage_key
            _assert(original_path.is_file(), "restored_original_missing")
            _assert(
                hashlib.sha256(original_path.read_bytes()).hexdigest() == resume.sha256,
                "restored_original_sha256_mismatch",
            )

        # Invoke the exact worker recovery function, then invoke its normal
        # claim function after the deliberate one-second retry backoff.  No
        # IMAP connection is opened: this proves a restart can reclaim a
        # durable queue record before slow external work begins.
        with database.session_factory() as session:
            mailbox_background_job_service._recover_expired_jobs(
                session,
                settings=settings,
                now=datetime.now(timezone.utc),
            )

        with database.session_factory() as session:
            set_organization_context(session, primary.id)
            try:
                recovered = session.get(MailboxBackgroundJob, expired_job.id)
            finally:
                clear_organization_context(session)
            _assert(recovered is not None, "recovered_job_missing")
            _assert(recovered.status == "queued", "expired_lease_not_requeued")
            _assert(
                recovered.last_error == "mailbox_background_job_lease_expired",
                "expired_lease_error_missing",
            )
            _assert(recovered.lease_owner is None and recovered.lease_expires_at is None, "expired_lease_not_cleared")
            _assert(recovered.next_attempt_at is not None, "expired_lease_retry_not_scheduled")
            wait_seconds = max(
                0.0,
                (recovered.next_attempt_at - datetime.now(timezone.utc)).total_seconds(),
            )

        if wait_seconds:
            time.sleep(wait_seconds + 0.15)
        claimed = mailbox_background_job_service._claim_next_job(
            database,
            settings=settings,
            worker_id="recovery-regression-worker",
        )
        _assert(claimed is not None, "recovered_job_not_claimable")
        _assert(claimed.organization_id == primary.id, "recovered_job_claimed_cross_workspace")

        with database.session_factory() as session:
            set_organization_context(session, primary.id)
            try:
                reclaimed = session.get(MailboxBackgroundJob, expired_job.id)
            finally:
                clear_organization_context(session)
            set_organization_context(session, secondary.id)
            try:
                untouched = session.scalar(
                    select(MailboxBackgroundJob).where(
                        MailboxBackgroundJob.status == "queued"
                    )
                )
            finally:
                clear_organization_context(session)
            _assert(reclaimed is not None, "reclaimed_job_missing")
            _assert(reclaimed.status == "running", "recovered_job_not_running_after_claim")
            _assert(reclaimed.lease_owner == "recovery-regression-worker", "recovered_job_owner_mismatch")
            _assert(untouched is not None, "secondary_workspace_job_missing")
            _assert(untouched.organization_id == secondary.id, "secondary_workspace_job_changed")
            _assert(untouched.attempt_count == 0, "secondary_workspace_job_claimed")
    finally:
        database.dispose()

    print("runtime-postgres-restore-and-lease-recovery: passed")


def main() -> None:
    parser = argparse.ArgumentParser(description="Run release regression checks inside the application image.")
    parser.add_argument(
        "mode",
        choices=("documents", "database-seed", "database-verify"),
        help="The isolated runtime check to execute.",
    )
    arguments = parser.parse_args()
    if arguments.mode == "documents":
        run_document_regression()
    elif arguments.mode == "database-seed":
        run_database_seed()
    else:
        run_database_verify()


if __name__ == "__main__":
    main()
